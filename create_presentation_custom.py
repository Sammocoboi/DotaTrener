from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

BG = RGBColor(0x0F, 0x0F, 0x1A)
CARD = RGBColor(0x1A, 0x1A, 0x2E)
ACCENT = RGBColor(0x6C, 0x5C, 0xE7)
ACCENT2 = RGBColor(0xE0, 0x45, 0x5B)
WHITE = RGBColor(0xF0, 0xF0, 0xF5)
GRAY = RGBColor(0x94, 0x94, 0xA8)
GREEN = RGBColor(0x2E, 0xCC, 0x71)
ORANGE = RGBColor(0xF3, 0x9C, 0x12)

def set_bg(slide, color=BG):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color

def shape(slide, l, t, w, h, color):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    s.fill.solid(); s.fill.fore_color.rgb = color; s.line.fill.background()
    return s

def text(slide, l, t, w, h, txt, sz=18, bold=False, color=WHITE, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = txt
    p.font.size = Pt(sz); p.font.bold = bold; p.font.color.rgb = color; p.font.name = 'Calibri'; p.alignment = align
    return tb

def bullets(slide, l, t, w, h, items, sz=14, color=GRAY):
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame; tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item; p.font.size = Pt(sz); p.font.color.rgb = color; p.font.name = 'Calibri'; p.space_after = Pt(4)
    return tb

def bar(slide, l, t, w=0.08, h=0.5, color=ACCENT):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = color; s.line.fill.background()

def card(slide, l, t, w, h, title, lines, title_sz=15, body_sz=12):
    shape(slide, l, t, w, h, CARD)
    bar(slide, l + Inches(0.02), t + Inches(0.15), 0.06, 0.35)
    text(slide, l + Inches(0.2), t + Inches(0.1), w - Inches(0.3), Inches(0.4), title, sz=title_sz, bold=True)
    bullets(slide, l + Inches(0.2), t + Inches(0.5), w - Inches(0.3), h - Inches(0.6), lines, sz=body_sz)

# ===== SLIDE 1 — Title =====
s = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s)
shape(s, Inches(0), Inches(0), prs.slide_width, prs.slide_height, RGBColor(0x15, 0x12, 0x2B))
shape(s, Inches(5.8), Inches(2.6), Inches(1.7), Inches(0.06), ACCENT)
text(s, Inches(1), Inches(1.0), Inches(11), Inches(1), 'DotaTrener', sz=64, bold=True, align=PP_ALIGN.CENTER)
text(s, Inches(1), Inches(1.9), Inches(11), Inches(0.7), 'Платформа для поиска тренеров по Dota 2', sz=26, color=GRAY, align=PP_ALIGN.CENTER)
text(s, Inches(1), Inches(3.0), Inches(11), Inches(0.5), 'Кроссплатформенное мобильное приложение (Android + Web)', sz=20, color=ACCENT, align=PP_ALIGN.CENTER)
text(s, Inches(1), Inches(3.6), Inches(11), Inches(0.6), 'React Native (Expo) + Django REST Framework', sz=16, color=GRAY, align=PP_ALIGN.CENTER)
text(s, Inches(1), Inches(4.8), Inches(11), Inches(0.5), 'Интеграция со Steam и STRATZ API  •  Встроенный мессенджер  •  Система услуг и отзывов', sz=14, color=GRAY, align=PP_ALIGN.CENTER)
text(s, Inches(1), Inches(6.5), Inches(11), Inches(0.5), '2025', sz=14, color=GRAY, align=PP_ALIGN.CENTER)

# ===== SLIDE 2 — Проблема =====
s = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s)
text(s, Inches(0.8), Inches(0.4), Inches(10), Inches(0.7), 'Проблема', sz=38, bold=True)
bar(s, Inches(0.8), Inches(1.0), 1.2, 0.05)

items = [
    ('Тысячи игроков хотят расти в рейтинге', 'но не знают, где найти хорошего тренера'),
    ('Разрозненные решения', 'Discord, Telegram, отдельные сайты — нет единой платформы'),
    ('Нет аналитики в одном месте', 'статистика матчей, герои, мета — всё разбросано'),
    ('Тренерам сложно', 'находить клиентов, управлять услугами и получать отзывы'),
]
for i, (a, b) in enumerate(items):
    y = 1.5 + i * 1.4
    shape(s, Inches(0.8), Inches(y), Inches(11.7), Inches(1.2), RGBColor(0x1E, 0x1E, 0x32))
    text(s, Inches(1.2), Inches(y + 0.1), Inches(11), Inches(0.4), a, sz=20, bold=True, color=ACCENT2)
    text(s, Inches(1.2), Inches(y + 0.55), Inches(11), Inches(0.4), b, sz=16, color=GRAY)

# ===== SLIDE 3 — Решение =====
s = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s)
text(s, Inches(0.8), Inches(0.4), Inches(10), Inches(0.7), 'Решение — DotaTrener', sz=38, bold=True)
bar(s, Inches(0.8), Inches(1.0), 1.2, 0.05)

text(s, Inches(0.8), Inches(1.4), Inches(11), Inches(0.4), 'Вся экосистема Dota 2 коучинга в одном приложении', sz=22, color=WHITE)

features = [
    ('Каталог тренеров', 'Фильтрация по рейтингу, цене, специализации'),
    ('Статистика матчей', 'Интеграция со Steam и STRATZ API: герои, KDA, GPM/XPM, графики'),
    ('Каталог героев', 'Мета, билды предметов и способностей, винрейт'),
    ('Мессенджер', 'Встроенный чат между игроком и тренером'),
    ('Система услуг', 'Тренеры создают предложения, игроки выбирают'),
    ('Отзывы и рейтинг', 'Оценки 1–5 с текстовыми комментариями'),
    ('Подписка', 'Доступ к эксклюзивным материалам и занятиям'),
    ('Панель администратора', 'Управление платформой и модерация'),
]
for i, (a, b) in enumerate(features):
    col, row = i % 4, i // 4
    x = 0.8 + col * 3.1; y = 2.0 + row * 2.6
    card(s, Inches(x), Inches(y), Inches(2.8), Inches(2.2), a, [b])

# ===== SLIDE 4 — Целевая аудитория =====
s = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s)
text(s, Inches(0.8), Inches(0.4), Inches(10), Inches(0.7), 'Целевая аудитория', sz=38, bold=True)
bar(s, Inches(0.8), Inches(1.0), 1.2, 0.05)

card(s, Inches(0.8), Inches(1.4), Inches(3.7), Inches(5.5),
    '🎮  Игроки (Player)', [
        'Ищут тренеров для прокачки скилла',
        'Хотят анализировать свою статистику',
        'Выбирают по рейтингу, цене, специализации',
        'Общаются с тренером через мессенджер',
        'Смотрят мету и билды героев',
    ])
card(s, Inches(4.8), Inches(1.4), Inches(3.7), Inches(5.5),
    '📢  Тренеры (Coach)', [
        'Предлагают услуги обучения игрокам',
        'Выкладывают видео-уроки (до 600 MB)',
        'Получают отзывы и рейтинг',
        'Создают эксклюзивные материалы по подписке',
        'Монетизируют свой опыт',
    ])
card(s, Inches(8.8), Inches(1.4), Inches(3.7), Inches(5.5),
    '⚙  Администраторы (Admin)', [
        'Модерируют заявки на становление тренером',
        'Управляют контентом и пользователями',
        'Контролируют качество услуг',
        'Работают через Django Admin',
    ])

# ===== SLIDE 5 — Ключевая идея =====
s = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s)
text(s, Inches(0.8), Inches(0.4), Inches(10), Inches(0.7), 'Ключевая идея', sz=38, bold=True)
bar(s, Inches(0.8), Inches(1.0), 1.2, 0.05)

shape(s, Inches(0.8), Inches(1.5), Inches(11.7), Inches(5.2), RGBColor(0x15, 0x15, 0x28))
text(s, Inches(1.5), Inches(1.8), Inches(10.5), Inches(0.8),
    'DotaTrener делает киберспортивный коучинг доступным, прозрачным и удобным',
    sz=26, bold=True, color=ACCENT, align=PP_ALIGN.CENTER)

parts = [
    ('Доступность', 'Вся экосистема в одном приложении: от анализа матчей до записи на занятие', ACCENT),
    ('Прозрачность', 'Реальные отзывы, верифицированный MMR и детальная статистика тренеров', GREEN),
    ('Удобство', 'Мобильное приложение + Web — занимайся где угодно', ORANGE),
    ('Эффективность', 'Современные инструменты аналитики для отслеживания прогресса', ACCENT2),
]
for i, (a, b, c) in enumerate(parts):
    x = 0.8 + i * 3.15
    shape(s, Inches(x), Inches(2.8), Inches(2.8), Inches(3.5), CARD)
    bar(s, Inches(x + 0.02), Inches(2.9), 0.06, 0.35, c)
    text(s, Inches(x + 0.2), Inches(2.9), Inches(2.4), Inches(0.4), a, sz=22, bold=True, color=c)
    text(s, Inches(x + 0.2), Inches(3.5), Inches(2.4), Inches(2.5), b, sz=14, color=GRAY)

# ===== SLIDE 6 — Технологический стек =====
s = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s)
text(s, Inches(0.8), Inches(0.4), Inches(10), Inches(0.7), 'Технологический стек', sz=38, bold=True)
bar(s, Inches(0.8), Inches(1.0), 1.2, 0.05)

shape(s, Inches(1.5), Inches(1.5), Inches(4.5), Inches(2.5), CARD)
bar(s, Inches(1.52), Inches(1.6), 0.06, 0.35, ACCENT)
text(s, Inches(1.8), Inches(1.5), Inches(4), Inches(0.4), 'Backend', sz=22, bold=True)
bullets(s, Inches(1.8), Inches(2.0), Inches(4), Inches(1.8), [
    'Python 3.10+ / Django 4.2 / DRF',
    'PostgreSQL / SQLite',
    'Аутентификация: Steam OpenID + Email',
    'Модели: Player, Coach, Match, Review, Chat',
    'Stratz GraphQL + OpenDota REST клиенты',
])

shape(s, Inches(7.3), Inches(1.5), Inches(4.5), Inches(2.5), CARD)
bar(s, Inches(7.32), Inches(1.6), 0.06, 0.35, ACCENT2)
text(s, Inches(7.6), Inches(1.5), Inches(4), Inches(0.4), 'Frontend', sz=22, bold=True)
bullets(s, Inches(7.6), Inches(2.0), Inches(4), Inches(1.8), [
    'React Native 0.85 + Expo 56',
    'TypeScript 6.0, 30+ экранов',
    'Expo Router (файловая навигация)',
    'Android + Web (React Native Web)',
    'SecureStore, Axios, темы оформления',
])

shape(s, Inches(1.5), Inches(4.4), Inches(10.3), Inches(2.5), CARD)
bar(s, Inches(1.52), Inches(4.5), 0.06, 0.35, ORANGE)
text(s, Inches(1.8), Inches(4.4), Inches(4), Inches(0.4), 'Внешние интеграции', sz=22, bold=True)
bullets(s, Inches(1.8), Inches(4.9), Inches(9.5), Inches(1.8), [
    'Stratz GraphQL API — матчи, герои, мета-статистика, билды предметов и способностей',
    'OpenDota REST API — резервный источник для матчей и статистики',
    'Steam Web API — профили, аватарки, OpenID аутентификация',
    'CDN для изображений героев и предметов',
])

# ===== SLIDE 7 — Функциональность детально =====
s = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s)
text(s, Inches(0.8), Inches(0.4), Inches(10), Inches(0.7), 'Функциональные блоки', sz=38, bold=True)
bar(s, Inches(0.8), Inches(1.0), 1.2, 0.05)

funcs = [
    ('Анализ матчей', GREEN, [
        'История 20 последних матчей',
        'KDA, GPM, XPM для каждого',
        'Все 10 игроков с героями',
        'Предметы, способности, таланты',
        'График networth, сравнение урона',
    ]),
    ('Каталог тренеров', ACCENT, [
        'Фильтр по рейтингу, цене, специализации',
        'Детальный профиль с био и MMR',
        'Создание/редактирование услуг',
        'Видео-уроки до 600 MB',
        'Отзывы и оценки 1–5 звёзд',
    ]),
    ('Герои и мета', ORANGE, [
        'Каталог всех героев с атрибутами',
        'Винрейт, пикрейт, банрейт',
        'Средние KDA/GPM/XPM по герою',
        'Билды предметов по фазам игры',
        'Билды способностей и таланты',
    ]),
    ('Коммуникация', ACCENT2, [
        'Встроенный мессенджер',
        'Диалог: игрок ↔ тренер',
        'Привязка к услуге',
        'Список диалогов с уведомлениями',
    ]),
]
for i, (a, c, ls) in enumerate(funcs):
    x = 0.8 + i * 3.15
    shape(s, Inches(x), Inches(1.4), Inches(2.8), Inches(5.5), CARD)
    bar(s, Inches(x + 0.02), Inches(1.5), 0.06, 0.35, c)
    text(s, Inches(x + 0.2), Inches(1.5), Inches(2.4), Inches(0.4), a, sz=18, bold=True, color=c)
    bullets(s, Inches(x + 0.2), Inches(2.0), Inches(2.4), Inches(4.5), ls, sz=13)

# ===== SLIDE 8 — Заключение =====
s = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s)
shape(s, Inches(0), Inches(0), prs.slide_width, prs.slide_height, RGBColor(0x15, 0x12, 0x2B))
shape(s, Inches(5.8), Inches(2.5), Inches(1.7), Inches(0.06), ACCENT)
text(s, Inches(1), Inches(1.2), Inches(11), Inches(1), 'DotaTrener', sz=56, bold=True, align=PP_ALIGN.CENTER)
text(s, Inches(1), Inches(2.0), Inches(11), Inches(0.6), 'Киберспортивный коучинг без границ', sz=24, color=GRAY, align=PP_ALIGN.CENTER)
text(s, Inches(1), Inches(3.2), Inches(11), Inches(1.5),
    'React Native (Expo) + Django REST Framework\n'
    'Stratz API  •  OpenDota API  •  Steam OpenID\n'
    'Android + Web  •  Русскоязычный интерфейс',
    sz=18, color=GRAY, align=PP_ALIGN.CENTER)
shape(s, Inches(3.5), Inches(5.0), Inches(6.3), Inches(0.04), ACCENT)
text(s, Inches(1), Inches(5.3), Inches(11), Inches(0.5), 'Спасибо за внимание!', sz=30, bold=True, color=ACCENT, align=PP_ALIGN.CENTER)
text(s, Inches(1), Inches(6.0), Inches(11), Inches(0.5), 'Вопросы?', sz=20, color=GRAY, align=PP_ALIGN.CENTER)

prs.save('C:\\Users\\Sammo Soboi\\Documents\\GitHub\\DotaTrener\\DotaTrener_Presentation_Custom.pptx')
print('Presentation saved successfully!')
