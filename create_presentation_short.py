from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

BG_DARK = RGBColor(0x0F, 0x0F, 0x1A)
BG_CARD = RGBColor(0x1A, 0x1A, 0x2E)
ACCENT = RGBColor(0x6C, 0x5C, 0xE7)
ACCENT2 = RGBColor(0xE0, 0x45, 0x5B)
TEXT_WHITE = RGBColor(0xF0, 0xF0, 0xF5)
TEXT_GRAY = RGBColor(0x94, 0x94, 0xA8)
GREEN = RGBColor(0x2E, 0xCC, 0x71)
ORANGE = RGBColor(0xF3, 0x9C, 0x12)
SCREENSHOT_BG = RGBColor(0x25, 0x25, 0x3A)

def set_bg(slide, color=BG_DARK):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_shape(slide, left, top, width, height, color, alpha=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape

def add_text(slide, left, top, width, height, text, size=18, bold=False, color=TEXT_WHITE, align=PP_ALIGN.LEFT, font_name='Calibri'):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = font_name
    p.alignment = align
    return txBox

def add_bullet_text(slide, left, top, width, height, items, size=16, color=TEXT_WHITE, spacing=Pt(8)):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = item
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.font.name = 'Calibri'
        p.space_after = spacing
        p.level = 0
    return txBox

def add_accent_bar(slide, left, top, width=Inches(0.08), height=Inches(0.5), color=ACCENT):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape

def add_card(slide, left, top, width, height, title, body_lines, title_size=16, body_size=13):
    card = add_shape(slide, left, top, width, height, BG_CARD)
    add_accent_bar(slide, left + Inches(0.02), top + Inches(0.15), Inches(0.06), Inches(0.35), ACCENT)
    add_text(slide, left + Inches(0.2), top + Inches(0.1), width - Inches(0.3), Inches(0.4), title, size=title_size, bold=True)
    txBox = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.5), width - Inches(0.3), height - Inches(0.6))
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, line in enumerate(body_lines):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = line
        p.font.size = Pt(body_size)
        p.font.color.rgb = TEXT_GRAY
        p.font.name = 'Calibri'
        p.space_after = Pt(4)
    return card

# ========== SLIDE 1: Title ==========
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)
add_shape(slide, Inches(0), Inches(0), prs.slide_width, prs.slide_height, RGBColor(0x15, 0x12, 0x2B))
add_shape(slide, Inches(5.8), Inches(2.8), Inches(1.7), Inches(0.06), ACCENT)
add_text(slide, Inches(1), Inches(1.2), Inches(11), Inches(1), 'DotaTrener', size=60, bold=True, align=PP_ALIGN.CENTER)
add_text(slide, Inches(1), Inches(2.0), Inches(11), Inches(0.8), 'Платформа для поиска тренеров и анализа матчей Dota 2', size=24, color=TEXT_GRAY, align=PP_ALIGN.CENTER)
add_text(slide, Inches(1), Inches(3.2), Inches(11), Inches(1), 'Fullstack-приложение: React Native (Expo) + Django REST Framework', size=18, color=ACCENT, align=PP_ALIGN.CENTER)
add_text(slide, Inches(1), Inches(4.5), Inches(11), Inches(1), 'Backend: Python / Django \u2022 Frontend: TypeScript / React Native\nStratz API \u2022 OpenDota API \u2022 Steam OpenID', size=16, color=TEXT_GRAY, align=PP_ALIGN.CENTER)

# ========== SLIDE 2: Problem & Solution ==========
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)
add_text(slide, Inches(0.8), Inches(0.4), Inches(10), Inches(0.7), 'О проекте', size=36, bold=True)
add_accent_bar(slide, Inches(0.8), Inches(1.0), Inches(1.2), Inches(0.05), ACCENT)

add_card(slide, Inches(0.8), Inches(1.4), Inches(5.5), Inches(2.8),
    'Проблема', [
        'Сложно найти квалифицированного тренера по Dota 2',
        'Нет единой платформы для анализа статистики и поиска тренера',
        'Разрозненные инструменты: Discord, Telegram, отдельные сайты',
        'Тренерам сложно находить клиентов и управлять услугами',
        'Нет встроенной аналитики матчей для оценки прогресса',
    ])

add_card(slide, Inches(6.8), Inches(1.4), Inches(5.5), Inches(2.8),
    'Решение — DotaTrener', [
        'Мобильное + Web-приложение для Dota 2 сообщества',
        'Маркетплейс услуг профессиональных тренеров',
        'Интеграция с API Dota 2 (Stratz, OpenDota, Steam)',
        'Аналитика матчей, героев, меты и билдов',
        'Встроенный мессенджер для игроков и тренеров',
        'Русскоязычный интерфейс и документация',
    ])

add_card(slide, Inches(0.8), Inches(4.6), Inches(11.5), Inches(2.2),
    'Ключевые возможности', [
        'Автоматический подбор тренера по рангу, специализации и цене',
        'Просмотр статистики матчей через Stratz и Steam API',
        'Внутренний мессенджер для общения игрока и тренера',
        'Управление услугами, видео-уроками и отзывами для тренеров',
        'Мета-статистика героев: винрейт, пикрейт, билды, предметы',
    ])

# ========== SLIDE 3: User Roles ==========
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)
add_text(slide, Inches(0.8), Inches(0.4), Inches(10), Inches(0.7), 'Роли пользователей', size=36, bold=True)
add_accent_bar(slide, Inches(0.8), Inches(1.0), Inches(1.2), Inches(0.05), ACCENT)

add_card(slide, Inches(0.8), Inches(1.4), Inches(3.7), Inches(5.2),
    'Игроки (Players)', [
        'Русскоязычные игроки Dota 2 любого уровня',
        'Хотят улучшить навыки через профессиональный коучинг',
        'Ищут удобный каталог тренеров с фильтрацией',
        'Нуждаются в анализе своей статистики и матчей',
        'Используют: история матчей, билды героев, поиск тренера',
    ])

add_card(slide, Inches(4.8), Inches(1.4), Inches(3.7), Inches(5.2),
    'Тренеры (Coaches)', [
        'Опытные игроки Divine/Immortal ранга',
        'Хотят находить клиентов и монетизировать опыт',
        'Управление услугами (цена, описание, длительность)',
        'Загрузка видео-уроков (до 600 MB)',
        'Отзывы и рейтинг, подписка на сервис',
    ])

add_card(slide, Inches(8.8), Inches(1.4), Inches(3.7), Inches(5.2),
    'Администраторы (Admins)', [
        'Модерация заявок на становление тренером',
        'Управление пользователями через Django Admin',
        'Контроль качества услуг и контента',
        'Просмотр и управление всей платформой',
        'Одобрение/отклонение заявок тренеров',
    ])

# ========== SLIDE 4: Key Features ==========
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)
add_text(slide, Inches(0.8), Inches(0.4), Inches(10), Inches(0.7), 'Функциональность', size=36, bold=True)
add_accent_bar(slide, Inches(0.8), Inches(1.0), Inches(1.2), Inches(0.05), ACCENT)

add_card(slide, Inches(0.8), Inches(1.4), Inches(3.7), Inches(2.8),
    'Анализ матчей', [
        'История последних 20 матчей игрока',
        'Детальный разбор: все 10 игроков, KDA, GPM, XPM',
        'Предметы, способности, таланты',
        'График нетворса и сравнение урона',
    ])

add_card(slide, Inches(4.8), Inches(1.4), Inches(3.7), Inches(2.8),
    'Каталог героев и мета', [
        'Все герои Dota 2 с атрибутами и ролями',
        'Мета-статистика: винрейт, пикрейт, банрейт',
        'Популярные билды предметов и способностей',
        'Средние KDA, GPM, XPM по каждому герою',
    ])

add_card(slide, Inches(8.8), Inches(1.4), Inches(3.7), Inches(2.8),
    'Тренерская платформа', [
        'Каталог тренеров с фильтрацией',
        'Профиль с рейтингом и верифицированным MMR',
        'Управление услугами и видео-уроками',
        'Система отзывов 1-5 звезд',
    ])

add_card(slide, Inches(0.8), Inches(4.6), Inches(3.7), Inches(2.2),
    'Мессенджер', [
        'Встроенный чат между игроком и тренером',
        'Привязка к конкретной услуге',
        'Список диалогов с непрочитанными',
    ])

add_card(slide, Inches(4.8), Inches(4.6), Inches(3.7), Inches(2.2),
    'Аутентификация', [
        'Вход через Steam OpenID',
        'Регистрация по Email + пароль',
        'Привязка/отвязка Steam аккаунта',
        'Безопасное хранение токенов',
    ])

add_card(slide, Inches(8.8), Inches(4.6), Inches(3.7), Inches(2.2),
    'Администрирование', [
        'Модерация заявок тренеров',
        'Django Admin панель',
        'Управление подписками',
        'Контроль качества услуг',
    ])

# ========== SLIDE 5: Architecture & Tech ==========
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)
add_text(slide, Inches(0.8), Inches(0.4), Inches(10), Inches(0.7), 'Архитектура и технологии', size=36, bold=True)
add_accent_bar(slide, Inches(0.8), Inches(1.0), Inches(1.2), Inches(0.05), ACCENT)

add_shape(slide, Inches(1.5), Inches(1.5), Inches(4.5), Inches(2.5), BG_CARD)
add_accent_bar(slide, Inches(1.52), Inches(1.6), Inches(0.06), Inches(0.35), ACCENT)
add_text(slide, Inches(1.8), Inches(1.5), Inches(4), Inches(0.4), 'Backend', size=20, bold=True)
add_bullet_text(slide, Inches(1.8), Inches(2.0), Inches(4), Inches(1.8), [
    'Django 4.2 + Django REST Framework',
    'PostgreSQL / SQLite',
    'Модели: Player, Coach, Match, Review',
    'Steam OpenID + Email авторизация',
], size=14, color=TEXT_GRAY)

add_shape(slide, Inches(7.3), Inches(1.5), Inches(4.5), Inches(2.5), BG_CARD)
add_accent_bar(slide, Inches(7.32), Inches(1.6), Inches(0.06), Inches(0.35), ACCENT2)
add_text(slide, Inches(7.6), Inches(1.5), Inches(4), Inches(0.4), 'Frontend', size=20, bold=True)
add_bullet_text(slide, Inches(7.6), Inches(2.0), Inches(4), Inches(1.8), [
    'React Native 0.85 + Expo 56',
    'TypeScript 6.0, 30+ экранов',
    'Expo Router (файловая навигация)',
    'Android + Web (React Native Web)',
], size=14, color=TEXT_GRAY)

add_shape(slide, Inches(1.5), Inches(4.4), Inches(10.3), Inches(2.5), BG_CARD)
add_accent_bar(slide, Inches(1.52), Inches(4.5), Inches(0.06), Inches(0.35), ORANGE)
add_text(slide, Inches(1.8), Inches(4.4), Inches(4), Inches(0.4), 'Внешние интеграции', size=20, bold=True)
add_bullet_text(slide, Inches(1.8), Inches(4.9), Inches(9.5), Inches(1.8), [
    'Stratz GraphQL API \u2014 матчи, герои, мета-статистика, билды, предметы, способности',
    'OpenDota API \u2014 резервный источник данных для матчей и статистики',
    'Steam Web API \u2014 профили пользователей, аватарки, OpenID аутентификация',
    'Синхронизация данных через management команды (sync_stratz_meta и др.)',
], size=14, color=TEXT_GRAY)

# ========== SLIDE 6: External Integrations ==========
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)
add_text(slide, Inches(0.8), Inches(0.4), Inches(10), Inches(0.7), 'Интеграция с внешними сервисами', size=36, bold=True)
add_accent_bar(slide, Inches(0.8), Inches(1.0), Inches(1.2), Inches(0.05), ACCENT)

add_card(slide, Inches(0.8), Inches(1.4), Inches(3.7), Inches(5.3),
    'Stratz GraphQL API', [
        'Основной источник данных о Dota 2',
        'Профиль игрока и матчи по steamAccountId',
        'Мета героев, билды предметов и способностей',
        'HeroMeta: win/pick/ban rate, средние KDA/GPM/XPM',
        'HeroBuild: последовательность предметов с винрейтом',
        'HeroEndgameLoadout: оптимальные 6 слотов',
        'Работа с персональным токеном пользователя',
        'Cloudscraper для обхода Cloudflare',
    ], body_size=12)

add_card(slide, Inches(4.8), Inches(1.4), Inches(3.7), Inches(5.3),
    'OpenDota REST API', [
        'Дополнительный (резервный) источник',
        'Данные о матчах и статистика героев',
        'Используется как fallback при недоступности Stratz',
        'REST-запросы через библиотеку requests',
        'Синхронизация меты через sync_opendota_meta',
        'Обработка таймаутов и ошибок',
    ], body_size=12)

add_card(slide, Inches(8.8), Inches(1.4), Inches(3.7), Inches(5.3),
    'Steam Web API', [
        'Аутентификация через Steam OpenID',
        'Верификация пользователей через Steam',
        'Получение профиля: никнейм, аватар, страна',
        'OpenID валидация с проверкой сигнатуры',
        'Привязка Steam ID к профилю игрока',
        'CDN для изображений героев и предметов',
    ], body_size=12)

# ========== SLIDE 7: Development & Deployment ==========
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)
add_text(slide, Inches(0.8), Inches(0.4), Inches(10), Inches(0.7), 'Разработка и развёртывание', size=36, bold=True)
add_accent_bar(slide, Inches(0.8), Inches(1.0), Inches(1.2), Inches(0.05), ACCENT)

add_card(slide, Inches(0.8), Inches(1.4), Inches(3.7), Inches(2.8),
    'Backend', [
        'Python 3.10+ / Django 4.2+ / DRF',
        'Gunicorn (WSGI) + Nginx',
        'PostgreSQL (prod) / SQLite (dev)',
        'pytest + pytest-django для тестов',
        'Настройка через .env',
    ])

add_card(slide, Inches(4.8), Inches(1.4), Inches(3.7), Inches(2.8),
    'Frontend', [
        'React Native / Expo 56',
        'EAS Build для сборки APK',
        'Expo Go для разработки',
        'Android API 21+',
        'Web: любой современный браузер',
    ])

add_card(slide, Inches(8.8), Inches(1.4), Inches(3.7), Inches(2.8),
    'Документация', [
        'ГОСТ-стандарт: 6 разделов',
        'Техническое задание',
        'Пояснительная записка',
        'Описание программы',
        'Руководство пользователя и программиста',
    ])

add_card(slide, Inches(0.8), Inches(4.6), Inches(5.5), Inches(2.2),
    'Процесс разработки', [
        'Git для контроля версий',
        'requirements.txt / requirements-dev.txt',
        'Seed-скрипты для тестовых данных',
        'Django Admin панель для управления',
        'Autofield для всех первичных ключей',
    ])

add_card(slide, Inches(6.8), Inches(4.6), Inches(5.5), Inches(2.2),
    'Платформы', [
        'Android (основная мобильная платформа)',
        'Web (через React Native Web)',
        'iOS (сконфигурировано в app.json)',
        'Сервер: Linux + PostgreSQL + Nginx',
        'Продакшн: 83.222.16.188',
    ])

# ========== SLIDE 8: Conclusion ==========
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)
add_shape(slide, Inches(0), Inches(0), prs.slide_width, prs.slide_height, RGBColor(0x15, 0x12, 0x2B))
add_shape(slide, Inches(5.8), Inches(3.0), Inches(1.7), Inches(0.06), ACCENT)
add_text(slide, Inches(1), Inches(1.5), Inches(11), Inches(1), 'DotaTrener', size=52, bold=True, align=PP_ALIGN.CENTER)
add_text(slide, Inches(1), Inches(2.3), Inches(11), Inches(0.6), 'Полноценная платформа для Dota 2 коучинга', size=22, color=TEXT_GRAY, align=PP_ALIGN.CENTER)
add_text(slide, Inches(1), Inches(3.4), Inches(11), Inches(1.5),
    'React Native (Expo) + Django REST Framework\n'
    'Stratz API \u2022 OpenDota API \u2022 Steam OpenID\n'
    'Android + Web \u2022 Русскоязычный интерфейс',
    size=18, color=TEXT_GRAY, align=PP_ALIGN.CENTER)

add_text(slide, Inches(1), Inches(6.5), Inches(11), Inches(0.5), 'Спасибо за внимание!', size=28, bold=True, color=ACCENT, align=PP_ALIGN.CENTER)

prs.save('C:\\Users\\Sammo Soboi\\Documents\\GitHub\\DotaTrener\\DotaTrener_Presentation_Short.pptx')
print('Presentation saved successfully!')
